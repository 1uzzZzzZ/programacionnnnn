from openpyxl import Workbook

# Crear un libro nuevo
wb = Workbook()

# --- Hoja 1: Ventas ---
ws1 = wb.active
ws1.title = "Ventas"
ws1.append(["Producto", "Cantidad", "PrecioUnitario"])
ws1.append(["Laptop", 2, 15000])
ws1.append(["Mouse", 10, 200])

# --- Hoja 2: Gastos ---
ws2 = wb.create_sheet(title="Gastos")
ws2.append(["Concepto", "Monto"])
ws2.append(["Luz", 500])
ws2.append(["Internet", 400])

# --- Hoja 3: Resumen ---
ws3 = wb.create_sheet(title="Resumen")
ws3.append(["Tipo", "Total"])
# Calculamos totales simples para el ejemplo
total_ventas = (2 * 15000) + (10 * 200)
total_gastos = 500 + 400
ws3.append(["Total Ventas", total_ventas])
ws3.append(["Total Gastos", total_gastos])

wb.save("Reporte_Semana15.xlsx")
print("Archivo 'Reporte_Semana15.xlsx' creado correctamente.")