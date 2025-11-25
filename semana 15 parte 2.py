import tkinter as tk
from tkinter import messagebox
from openpyxl import load_workbook
import os

NOMBRE_ARCHIVO = "Reporte_Semana15.xlsx"

def guardar_datos():
    nombre = entry_nombre.get()
    calificacion = entry_calificacion.get()

    if not nombre or not calificacion:
        messagebox.showwarning("Error", "Por favor llena ambos campos")
        return

    try:
        # Cargar el archivo Excel existente
        wb = load_workbook(NOMBRE_ARCHIVO)
        
        # Buscar si existe hoja de Calificaciones, si no, crearla
        if "Calificaciones" in wb.sheetnames:
            ws = wb["Calificaciones"]
        else:
            ws = wb.create_sheet("Calificaciones")
            ws.append(["Nombre", "Calificación"]) # Encabezados

        # Guardar datos
        ws.append([nombre, float(calificacion)])
        wb.save(NOMBRE_ARCHIVO)
        
        messagebox.showinfo("Éxito", "Datos guardados en Excel")
        entry_nombre.delete(0, tk.END)
        entry_calificacion.delete(0, tk.END)

    except FileNotFoundError:
        messagebox.showerror("Error", f"No se encontró el archivo {NOMBRE_ARCHIVO}. Ejecuta el script 15.1 primero.")
    except ValueError:
        messagebox.showerror("Error", "La calificación debe ser un número.")

# Configuración de la ventana (Tkinter)
ventana = tk.Tk()
ventana.title("Semana 15.2: Ingreso de Datos")
ventana.geometry("300x200")

tk.Label(ventana, text="Nombre del Estudiante:").pack(pady=5)
entry_nombre = tk.Entry(ventana)
entry_nombre.pack(pady=5)

tk.Label(ventana, text="Calificación:").pack(pady=5)
entry_calificacion = tk.Entry(ventana)
entry_calificacion.pack(pady=5)

btn_guardar = tk.Button(ventana, text="Guardar en Excel", command=guardar_datos, bg="lightblue")
btn_guardar.pack(pady=20)

ventana.mainloop()